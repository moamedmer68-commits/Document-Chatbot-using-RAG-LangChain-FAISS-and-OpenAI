import os
import io
import hashlib
from typing import List, Tuple

import pandas as pd
import streamlit as st
from docx import Document as DocxDocument
from pypdf import PdfReader

from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_community.vectorstores import FAISS


SUPPORTED_TYPES = ["pdf", "csv", "docx", "txt", "md", "xlsx"]


st.set_page_config(page_title="RAG Assistant", page_icon="📚", layout="wide")


def resolve_openai_api_key() -> str:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if api_key:
        return api_key

    try:
        secret_key = st.secrets.get("OPENAI_API_KEY", "")
        if secret_key:
            os.environ["OPENAI_API_KEY"] = secret_key
            return secret_key
    except Exception:
        pass

    try:
        from google.colab import userdata  # type: ignore

        secret_key = userdata.get("OPENAI_API_KEY")
        if secret_key:
            os.environ["OPENAI_API_KEY"] = secret_key
            return secret_key
    except Exception:
        pass

    return ""


def file_digest(uploaded_files) -> str:
    h = hashlib.sha256()
    for f in uploaded_files:
        h.update(f.name.encode("utf-8"))
        h.update(f.getvalue())
    return h.hexdigest()


def serialize_documents(docs: List[Document]) -> Tuple[Tuple[str, tuple], ...]:
    serialized = []
    for d in docs:
        serialized.append((d.page_content, tuple(sorted(d.metadata.items()))))
    return tuple(serialized)


def format_docs(docs: List[Document]) -> str:
    formatted = []
    for doc in docs:
        source = doc.metadata.get("source", "unknown")
        page = doc.metadata.get("page")
        row = doc.metadata.get("row")
        sheet = doc.metadata.get("sheet")
        location_parts = []
        if sheet is not None:
            location_parts.append(f"sheet {sheet}")
        if page is not None:
            location_parts.append(f"page {page}")
        if row is not None:
            location_parts.append(f"row {row}")
        location = f" ({', '.join(location_parts)})" if location_parts else ""
        formatted.append(f"[Source: {source}{location}]\n{doc.page_content}")
    return "\n\n".join(formatted)


def load_pdf(uploaded_file) -> List[Document]:
    docs: List[Document] = []
    reader = PdfReader(io.BytesIO(uploaded_file.getvalue()))
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": uploaded_file.name,
                        "type": "pdf",
                        "page": i,
                    },
                )
            )
    return docs


def load_docx(uploaded_file) -> List[Document]:
    docs: List[Document] = []
    doc = DocxDocument(io.BytesIO(uploaded_file.getvalue()))

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if paragraphs:
        docs.append(
            Document(
                page_content="\n".join(paragraphs),
                metadata={"source": uploaded_file.name, "type": "docx", "section": "paragraphs"},
            )
        )

    for t_idx, table in enumerate(doc.tables, start=1):
        table_rows = []
        for r_idx, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                table_rows.append(f"row {r_idx}: " + " | ".join(values))
        if table_rows:
            docs.append(
                Document(
                    page_content="\n".join(table_rows),
                    metadata={
                        "source": uploaded_file.name,
                        "type": "docx",
                        "section": f"table_{t_idx}",
                    },
                )
            )
    return docs


def load_csv(uploaded_file) -> List[Document]:
    docs: List[Document] = []
    df = pd.read_csv(io.BytesIO(uploaded_file.getvalue()))
    df = df.fillna("")
    for idx, row in df.iterrows():
        row_text = "\n".join([f"{col}: {row[col]}" for col in df.columns])
        docs.append(
            Document(
                page_content=row_text,
                metadata={
                    "source": uploaded_file.name,
                    "type": "csv",
                    "row": int(idx) + 1,
                },
            )
        )
    return docs


def load_xlsx(uploaded_file) -> List[Document]:
    docs: List[Document] = []
    excel_bytes = io.BytesIO(uploaded_file.getvalue())
    xls = pd.ExcelFile(excel_bytes)

    for sheet_name in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet_name).fillna("")
        for idx, row in df.iterrows():
            row_text = "\n".join([f"{col}: {row[col]}" for col in df.columns])
            docs.append(
                Document(
                    page_content=row_text,
                    metadata={
                        "source": uploaded_file.name,
                        "type": "xlsx",
                        "sheet": sheet_name,
                        "row": int(idx) + 1,
                    },
                )
            )
    return docs


def load_text(uploaded_file) -> List[Document]:
    text = uploaded_file.getvalue().decode("utf-8", errors="ignore").strip()
    if not text:
        return []
    return [
        Document(
            page_content=text,
            metadata={"source": uploaded_file.name, "type": "text"},
        )
    ]


def load_uploaded_documents(uploaded_files) -> List[Document]:
    docs: List[Document] = []
    for uploaded_file in uploaded_files:
        ext = uploaded_file.name.rsplit(".", 1)[-1].lower()
        try:
            if ext == "pdf":
                docs.extend(load_pdf(uploaded_file))
            elif ext == "docx":
                docs.extend(load_docx(uploaded_file))
            elif ext == "csv":
                docs.extend(load_csv(uploaded_file))
            elif ext == "xlsx":
                docs.extend(load_xlsx(uploaded_file))
            elif ext in {"txt", "md"}:
                docs.extend(load_text(uploaded_file))
            else:
                st.warning(f"Unsupported file skipped: {uploaded_file.name}")
        except Exception as e:
            st.warning(f"Could not read {uploaded_file.name}: {e}")
    return docs


def split_documents(documents: List[Document]) -> List[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1200,
        chunk_overlap=200,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return splitter.split_documents(documents)


@st.cache_resource(show_spinner=False)
def build_vectorstore(_signature: str, documents: Tuple[Tuple[str, tuple], ...]):
    docs = [Document(page_content=pc, metadata=dict(md)) for pc, md in documents]
    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
    return FAISS.from_documents(docs, embeddings)


def build_rag_chain(vectorstore):
    retriever = vectorstore.as_retriever(search_kwargs={"k": 4})
    llm = ChatOpenAI(model="gpt-3.5-turbo", temperature=0)

    prompt = ChatPromptTemplate.from_template(
        """You are a precise RAG assistant.
Use only the retrieved context below to answer the user's question.
If the answer is not in the context, say clearly that the answer was not found in the uploaded files.
Prefer a direct, concise, helpful answer.
Answer in Arabic unless the user's question is in another language.

Question:
{question}

Retrieved context:
{context}

Answer:"""
    )

    chain = (
        {
            "context": retriever | format_docs,
            "question": RunnablePassthrough(),
        }
        | prompt
        | llm
        | StrOutputParser()
    )
    return chain, retriever


def init_state():
    defaults = {
        "messages": [],
        "vectorstore": None,
        "chain": None,
        "retriever": None,
        "docs_signature": None,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def reset_chat():
    st.session_state.messages = []


def main():
    init_state()
    api_key = resolve_openai_api_key()

    st.title("📚 RAG Assistant")
    st.caption("Upload files, build a knowledge base, then chat with your data.")

    with st.sidebar:
        st.header("Settings")
        if api_key:
            st.success("OpenAI API key loaded")
        else:
            st.error("No OPENAI_API_KEY found. Add it to Colab Secrets, Streamlit secrets, or environment variables.")

        uploaded_files = st.file_uploader(
            "Upload files",
            type=SUPPORTED_TYPES,
            accept_multiple_files=True,
            help="Supported: PDF, CSV, DOCX, XLSX, TXT, MD",
        )

        if st.button("Build / Refresh Knowledge Base", use_container_width=True):
            if not api_key:
                st.stop()
            if not uploaded_files:
                st.warning("Upload at least one file first.")
                st.stop()

            with st.spinner("Reading files and building vector index..."):
                raw_docs = load_uploaded_documents(uploaded_files)
                if not raw_docs:
                    st.error("No readable text was found in the uploaded files.")
                    st.stop()

                chunks = split_documents(raw_docs)
                signature = file_digest(uploaded_files)
                serialized_docs = serialize_documents(chunks)
                vectorstore = build_vectorstore(signature, serialized_docs)
                chain, retriever = build_rag_chain(vectorstore)

                st.session_state.vectorstore = vectorstore
                st.session_state.chain = chain
                st.session_state.retriever = retriever
                st.session_state.docs_signature = signature
                st.session_state.messages = []

            st.success(f"Knowledge base ready. Indexed {len(chunks)} chunks.")

        if st.button("Clear chat", use_container_width=True):
            reset_chat()
            st.rerun()

        st.markdown("---")
        st.markdown("**Tips**")
        st.markdown("- Upload more than one file if you want a shared knowledge base.")
        st.markdown("- CSV/XLSX are indexed row by row.")
        st.markdown("- PDF extraction works best on text-based PDFs.")

    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])
            if msg.get("sources"):
                with st.expander("Sources"):
                    for src in msg["sources"]:
                        st.write(src)

    user_question = st.chat_input("Ask a question about your uploaded files...")

    if user_question:
        if st.session_state.chain is None or st.session_state.retriever is None:
            st.warning("Upload files and click 'Build / Refresh Knowledge Base' first.")
            st.stop()

        st.session_state.messages.append({"role": "user", "content": user_question})
        with st.chat_message("user"):
            st.markdown(user_question)

        with st.chat_message("assistant"):
            with st.spinner("Thinking..."):
                answer = st.session_state.chain.invoke(user_question)
                retrieved_docs = st.session_state.retriever.invoke(user_question)

                sources = []
                for d in retrieved_docs:
                    source = d.metadata.get("source", "unknown")
                    page = d.metadata.get("page")
                    row = d.metadata.get("row")
                    sheet = d.metadata.get("sheet")
                    section = d.metadata.get("section")

                    details = []
                    if page is not None:
                        details.append(f"page {page}")
                    if sheet is not None:
                        details.append(f"sheet {sheet}")
                    if row is not None:
                        details.append(f"row {row}")
                    if section is not None:
                        details.append(str(section))

                    label = f"{source} — {', '.join(details)}" if details else source
                    sources.append(label)

                sources = list(dict.fromkeys(sources))

            st.markdown(answer)
            with st.expander("Sources"):
                for src in sources:
                    st.write(src)

        st.session_state.messages.append(
            {"role": "assistant", "content": answer, "sources": sources}
        )


if __name__ == "__main__":
    main()
